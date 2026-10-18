"""Convert line-oriented UTF-8 text with the existing filing style contract.

Only explicitly delimited pipe/aligned tables can be removed. Inspect first
when deletion boundaries are uncertain; this is not an OCR reconstruction tool.
"""
import argparse
import json
import re
from pathlib import Path

from docx import Document
import filing_docx as F


PAGE = re.compile(r"^\s*=+\s*第\s*\d+\s*页\s*=+\s*$")
SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}:?(?:\s*\|\s*:?-{3,}:?)+\s*\|?\s*$")
ALIGNED_SEPARATOR = re.compile(r"^\s*-{3,}(?:\s+-{3,})+\s*$")
CN = "一二三四五六七八九十百零〇两"


def is_row(line, kind):
    text = line.strip()
    if not text or PAGE.fullmatch(text):
        return False
    if kind == "pipe":
        return "|" in text
    return len(re.split(r"\s{2,}|\t+", text)) > 1


def table_ranges(lines):
    """Return half-open ranges anchored on explicit table separator rows.

    A prose paragraph with multiple spaces alone is never a deletion anchor.
    Blank lines, page markers and non-row text end a detected table.
    """
    ranges = []
    for i, line in enumerate(lines):
        kind = ("pipe" if SEPARATOR.fullmatch(line) else
                "aligned" if ALIGNED_SEPARATOR.fullmatch(line) else None)
        if kind is None or i == 0 or not is_row(lines[i - 1], kind):
            continue
        start, end = i - 1, i + 1
        while end < len(lines) and is_row(lines[end], kind):
            end += 1
        if ranges and start <= ranges[-1][1]:
            ranges[-1][1] = max(end, ranges[-1][1])
        else:
            ranges.append([start, end])
    return ranges


def style_for(text, profile):
    prospectus = profile == "招股说明书"
    if re.match(r"^第[" + CN + r"\d]+节\s*\S", text):
        return "Heading1" if prospectus else "001"
    if re.match(r"^[" + CN + r"]+、", text):
        return "Heading2" if prospectus else "002"
    if re.match(r"^[（(][" + CN + r"]+[）)]", text):
        return "Heading3" if prospectus else "003"
    # Avoid treating decimal numbers such as 1.25 as numbered headings.
    if re.match(r"^\d+(?:、|[.．](?!\d))", text):
        return "Heading4" if prospectus else "004"
    if re.match(r"^[（(]\d+[）)]", text):
        return "Heading5" if prospectus else "011"
    if text.startswith("单位："):
        return "TableUnit"
    if re.match(r"^注\d*[：:]", text):
        return "TableNote"
    return "BodyIndent" if prospectus else "005"


def convert(src, dst, profile="申报文件", drop_tables=False,
            inspect=False, header=None, footer_prefix=""):
    src = Path(src)
    if not inspect and src.resolve() == Path(dst).resolve():
        raise ValueError("输入与输出必须使用不同路径")
    # Decode strictly: do not silently lose source characters.
    lines = src.read_text(encoding="utf-8-sig").splitlines()
    ranges = table_ranges(lines)
    excluded = {i for start, end in ranges for i in range(start, end)} if drop_tables else set()
    kept = [line.strip() for i, line in enumerate(lines)
            if i not in excluded and line.strip() and not PAGE.fullmatch(line)]
    stats = {
        "source_lines": len(lines), "detected_tables": len(ranges),
        "table_line_ranges": [[a + 1, b] for a, b in ranges],
        "removed_table_lines": len(excluded),
        "removed_page_markers": sum(bool(PAGE.fullmatch(s)) for s in lines),
        "kept_paragraphs": len(kept), "profile": profile,
        "table_detection": "仅识别有分隔线的表；其他内容保留为文本",
    }
    if inspect:
        stats["sample"] = kept[:8]
        return stats, True
    if not kept:
        raise ValueError("没有可输出的正文；请检查输入及删表范围")
    doc = F.new_document(profile=profile, auto_numbering=False)
    for text in kept:
        doc.add_paragraph(text, style=F.get_style(doc, style_for(text, profile)))
    F.set_header(doc, doc.sections[0], header if header is not None else src.stem)
    F.set_footer(doc, doc.sections[0], footer_prefix)
    doc.save(str(dst))
    # Compare every retained text unit after serialization, in order.
    actual = [p.text for p in Document(str(dst)).paragraphs if p.text]
    stats["content_roundtrip_ok"] = actual == kept
    ok, findings = F.check(str(dst), profile=profile)
    stats["format_ok"] = ok
    stats["failures"] = sum(f["level"] == "FAIL" for f in findings)
    stats["warnings"] = sum(f["level"] == "WARN" for f in findings)
    stats["output"] = str(dst)
    return stats, ok and stats["content_roundtrip_ok"]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("src")
    parser.add_argument("dst", nargs="?")
    parser.add_argument("--profile", choices=["申报文件", "招股说明书"], default="申报文件")
    parser.add_argument("--drop-tables", action="store_true", help="仅用户要求删表时使用")
    parser.add_argument("--inspect", action="store_true", help="只查看识别范围，不写 Word")
    parser.add_argument("--header")
    parser.add_argument("--footer-prefix", default="")
    args = parser.parse_args()
    if not args.inspect and args.dst is None:
        parser.error("生成 Word 需要指定输出路径")
    try:
        stats, ok = convert(**vars(args))
    except (OSError, UnicodeError, ValueError) as exc:
        parser.exit(2, f"转换失败：{exc}\n")
    print(json.dumps(stats, ensure_ascii=False, indent=2))
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
